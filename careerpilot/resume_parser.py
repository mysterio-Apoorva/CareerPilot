import io
import re
from pydantic import BaseModel, Field

# Section labels route literal resume text; they are not allowed skill values.
HEADINGS = {
    "skills": ["skills", "technical skills", "technical expertise", "core competencies", "competencies",
               "programming languages", "languages", "frameworks", "libraries", "tools", "technologies",
               "databases", "tools and technologies", "tools & technologies"],
    "education": ["education", "academic background", "academic qualifications", "qualifications"],
    "projects": ["projects", "project", "personal projects", "academic projects", "selected projects", "project experience"],
    "experience": ["experience", "work experience", "professional experience", "internships", "research experience", "employment"],
    "certifications": ["certifications", "courses", "certificates", "courses and certifications", "certifications and courses"],
    "achievements": ["achievements", "awards", "honors", "leadership", "activities", "positions of responsibility", "competitions"],
    "other": ["summary", "professional summary", "objective", "profile", "contact", "interests", "references"],
}


class ResumeExtraction(BaseModel):
    education: str = ""
    academic_stage: str = ""
    graduation_year: int | None = None
    skills: list[str] = Field(default_factory=list)
    projects: str = ""
    experience: str = ""
    certifications: str = ""
    achievements: str = ""
    evidence: dict[str, list[str]] = Field(default_factory=dict)
    inferred_skills: list[str] = Field(default_factory=list)
    method: str = "Local section extraction; review required"


def _sections(text):
    sections = {name: [] for name in HEADINGS}
    active = "other"
    for original in text.splitlines():
        line = original.strip().strip("•●▪ ")
        if not line:
            continue
        matched = False
        for section, aliases in HEADINGS.items():
            for alias in sorted(aliases, key=len, reverse=True):
                match = re.match(r"^" + re.escape(alias) + r"\s*(?::|[–—])\s*(.*)$", line, re.I)
                if line.casefold() == alias or match:
                    active = section
                    if match and match.group(1).strip():
                        sections[active].append(match.group(1).strip())
                    matched = True
                    break
            if matched:
                break
        if not matched:
            sections[active].append(line)
    return sections


def _listed_skills(lines):
    found = []
    for line in lines:
        # Preserve arbitrary technologies exactly as supplied, including C++ and C#.
        values = line.split(":", 1)[-1]
        for token in re.split(r"[,;|•●▪\t]|\s/\s", values):
            token = token.strip().strip("- ")
            if token and token.casefold() not in {s.casefold() for s in found}:
                found.append(token)
    return found


def parse_profile(text):
    sections = _sections(text)
    parsed = ResumeExtraction()
    for field in ("education", "projects", "experience", "certifications", "achievements"):
        lines = sections[field]
        setattr(parsed, field, "\n".join(lines))
        if lines:
            parsed.evidence[field] = lines
    if not parsed.education:
        lines = [line.strip() for line in text.splitlines() if re.search(r"\b(?:B\.?\s?Tech|M\.?\s?Tech|B\.?Sc|M\.?Sc|Bachelor|Master|Ph\.?D|MBA)\b", line, re.I)]
        parsed.education = "\n".join(lines)
        if lines:
            parsed.evidence["education"] = lines
    parsed.skills = _listed_skills(sections["skills"])
    if parsed.skills:
        parsed.evidence["skills"] = sections["skills"]
    stage = re.search(r"[^\n]*(?:(?:first|second|third|fourth|final|1st|2nd|3rd|4th)\s*[- ]?year|semester\s*\d+)[^\n]*", text, re.I)
    if stage:
        parsed.academic_stage = stage.group().strip()
        parsed.evidence["academic_stage"] = [stage.group().strip()]
    # No date arithmetic or guessing from an education date range.
    graduation = re.search(r"[^\n]*(?:expected\s+(?:graduation|completion)|graduation(?:\s+date|\s+year)?|graduating|graduate(?:s|d)?\s+in)\s*[:–—-]?\s*(?:[A-Za-z]+\s+)?(20\d{2})[^\n]*", text, re.I)
    if graduation:
        parsed.graduation_year = int(graduation.group(1))
        parsed.evidence["graduation_year"] = [graduation.group().strip()]
    return parsed


def validate_gpt_profile(data, text):
    """Only resume-backed facts may populate a profile. Suggestions stay separate."""
    parsed = ResumeExtraction(method="GPT extraction with source-excerpt validation; review required")
    for field in ("education", "academic_stage", "projects", "experience", "certifications", "achievements"):
        records = data.get(field) or []
        if isinstance(records, dict):
            records = [records]
        if not isinstance(records, list):
            continue
        quotes = [r.get("quote", "").strip() for r in records if isinstance(r, dict)]
        quotes = list(dict.fromkeys(q for q in quotes if q and q in text))
        setattr(parsed, field, "\n".join(quotes))
        if quotes:
            parsed.evidence[field] = quotes
    for record in data.get("skills") or []:
        if not isinstance(record, dict):
            continue
        name, quote = str(record.get("name", "")).strip(), str(record.get("quote", "")).strip()
        if name and quote and quote in text and re.search(r"(?<!\w)" + re.escape(name) + r"(?!\w)", quote, re.I):
            if name.casefold() not in {s.casefold() for s in parsed.skills}:
                parsed.skills.append(name)
                parsed.evidence.setdefault("skills", []).append(quote)
    graduation = data.get("graduation_year")
    if isinstance(graduation, dict):
        value, quote = graduation.get("value"), str(graduation.get("quote", ""))
        if type(value) is int and 2000 <= value <= 2100 and quote and quote in text and str(value) in quote and re.search(r"graduat|expected|completion", quote, re.I):
            parsed.graduation_year = value
            parsed.evidence["graduation_year"] = [quote]
    parsed.inferred_skills = [str(s) for s in data.get("inferred_skills", []) if isinstance(s, str)]
    return parsed


def apply_extraction(profile, parsed):
    """Replace resume facts, retaining only separately chosen career preferences."""
    fields = parsed.model_dump(exclude={"evidence", "inferred_skills", "method"})
    return profile.model_copy(update=fields)


def extract_resume(data: bytes, filename: str):
    if len(data) > 5 * 1024 * 1024:
        raise ValueError("Please upload a file smaller than 5 MB.")
    try:
        if filename.lower().endswith(".pdf"):
            from pypdf import PdfReader
            if not data.startswith(b"%PDF"):
                raise ValueError("Not a valid PDF.")
            reader = PdfReader(io.BytesIO(data))
            if len(reader.pages) > 30:
                raise ValueError("Resume exceeds the 30-page limit.")
            text = "\n".join(p.extract_text() or "" for p in reader.pages)
        elif filename.lower().endswith(".docx"):
            from docx import Document
            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
            text += "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        else:
            raise ValueError("Upload a PDF or DOCX resume.")
    except Exception as exc:
        raise ValueError("Could not read this resume. Use the manual profile fields below.") from exc
    if len(text.strip()) < 30:
        raise ValueError("No readable text found. Scanned PDFs need OCR; use manual entry for this MVP.")
    if len(text) > 100000:
        raise ValueError("Resume exceeds the text limit; upload a shorter document or use manual entry.")
    return text


def explicit_skills(text):
    return _listed_skills(_sections(text)["skills"])
